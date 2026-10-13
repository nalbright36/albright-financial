"""Tests for LedgerExpense (business-wide expenses, not tied to either
ledger), the Expenses tab, the Overview tab's expense-inclusive tiles/
rows/monthly breakdown, the import_personal_ledger management command,
and confirmation that expenses never touch the Reselling-only metrics
(dashboard tiles, Scorecard, Personal Items, net_cash_flow's own
resale-only default)."""
import os
import tempfile
from datetime import date
from decimal import Decimal
from io import StringIO

import openpyxl
from django.contrib.auth.models import User
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from albright_reselling_app import ledger_metrics
from albright_reselling_app.management.commands.import_personal_ledger import (
    _find_existing_sale, _plan_selling_sheet,
)
from albright_reselling_app.models import LedgerEntry, LedgerExpense, LedgerSale


def _make_entry(owner, item="Test Item", cost="0.00", status="holding", ledger_type="resale", **extra):
    entry = LedgerEntry.objects.create(
        owner=owner, item=item, cost=Decimal(cost), status=status, ledger_type=ledger_type,
        purchase_date=timezone.localdate(), **extra,
    )
    entry.refresh_from_db()
    return entry


def _make_expense(owner, description="Supplies Item", amount="10.00", date_=None, category="supplies", **extra):
    return LedgerExpense.objects.create(
        owner=owner, description=description, amount=Decimal(amount), date=date_, category=category, **extra,
    )


# ---------------------------------------------------------------------------
# Model
# ---------------------------------------------------------------------------

class LedgerExpenseModelTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="tester", password="pw-not-real-12345")

    def test_default_category_is_supplies(self):
        expense = LedgerExpense.objects.create(owner=self.user, description="Tape", amount=Decimal("5.00"))
        self.assertEqual(expense.category, "supplies")

    def test_date_is_optional(self):
        expense = _make_expense(self.user, date_=None)
        self.assertIsNone(expense.date)

    def test_no_ledger_type_field(self):
        field_names = [f.name for f in LedgerExpense._meta.get_fields()]
        self.assertNotIn("ledger_type", field_names)


# ---------------------------------------------------------------------------
# Overview tab
# ---------------------------------------------------------------------------

class OverviewExpensesTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="tester", password="pw-not-real-12345")
        self.client.login(username="tester", password="pw-not-real-12345")

    def test_tiles_revenue_item_costs_expenses_net(self):
        today = timezone.localdate()
        _make_entry(self.user, item="Resale Sold", cost="10.00", sold_for="40.00", status="sold_out",
                    sold_date=today)
        personal = _make_entry(self.user, item="Personal Sold", status="sold_out", ledger_type="personal")
        LedgerSale.objects.create(entry=personal, sale_date=today, sale_price=Decimal("20.00"))
        _make_expense(self.user, amount="15.00", date_=today)

        context = ledger_metrics.overview_context(self.user)

        self.assertEqual(context["tiles"]["total_revenue"], Decimal("60.00"))  # 40 + 20
        self.assertEqual(context["tiles"]["item_costs"], Decimal("10.00"))  # personal's own cost is always 0
        self.assertEqual(context["tiles"]["expenses"], Decimal("15.00"))
        self.assertEqual(context["tiles"]["net"], Decimal("35.00"))  # 60 - 10 - 15

    def test_rows_include_expenses_row_costs_only(self):
        _make_expense(self.user, amount="30.00")

        context = ledger_metrics.overview_context(self.user)

        expense_row = next(r for r in context["rows"] if r["label"] == "Expenses")
        self.assertEqual(expense_row["costs"], Decimal("30.00"))
        self.assertIsNone(expense_row["revenue"])
        self.assertIsNone(expense_row["net"])

    def test_total_row_subtracts_expenses(self):
        today = timezone.localdate()
        _make_entry(self.user, item="Resale Sold", cost="10.00", sold_for="40.00", status="sold_out",
                    sold_date=today)
        _make_expense(self.user, amount="5.00", date_=today)

        total_row = ledger_metrics.overview_context(self.user)["total_row"]

        self.assertEqual(total_row["revenue"], Decimal("40.00"))
        self.assertEqual(total_row["costs"], Decimal("15.00"))  # 10 item cost + 5 expense
        self.assertEqual(total_row["net"], Decimal("25.00"))  # 40 - 15

    def test_monthly_breakdown_includes_dated_expense_in_its_own_month(self):
        now = timezone.now()
        this_month_start = timezone.localdate().replace(day=1)
        _make_expense(self.user, amount="12.00", date_=this_month_start)

        context = ledger_metrics.overview_context(self.user, now=now)

        this_month_row = next(
            r for r in context["monthly_rows"] if r["month"].month == now.month and r["month"].year == now.year
        )
        self.assertEqual(this_month_row["expenses"], Decimal("12.00"))
        self.assertEqual(this_month_row["combined_net"], Decimal("-12.00"))

    def test_undated_expense_excluded_from_monthly_but_included_all_time(self):
        _make_expense(self.user, amount="50.00", date_=None)

        context = ledger_metrics.overview_context(self.user)

        self.assertEqual(sum((r["expenses"] for r in context["monthly_rows"]), Decimal("0")), Decimal("0"))
        self.assertEqual(context["tiles"]["expenses"], Decimal("50.00"))
        self.assertEqual(context["undated_expenses_total"], Decimal("50.00"))

    def test_overview_page_renders_expense_tiles_and_note(self):
        _make_expense(self.user, description="Tape", amount="9.99", date_=None)

        response = self.client.get(reverse("albright_reselling_app:ledger_overview"))

        self.assertContains(response, "Item Costs")
        self.assertContains(response, "Expenses")
        self.assertContains(response, "$9.99")
        self.assertContains(response, "undated expenses")


# ---------------------------------------------------------------------------
# Expenses tab
# ---------------------------------------------------------------------------

class ExpensesTabTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="tester", password="pw-not-real-12345")
        self.client.login(username="tester", password="pw-not-real-12345")

    def test_expense_appears_only_on_expenses_tab_and_overview(self):
        _make_expense(self.user, description="Unique Grading Fee", amount="25.00", category="grading")

        expenses_resp = self.client.get(reverse("albright_reselling_app:ledger_expenses"))
        overview_resp = self.client.get(reverse("albright_reselling_app:ledger_overview"))
        resale_resp = self.client.get(reverse("albright_reselling_app:ledger_resale"))
        personal_resp = self.client.get(reverse("albright_reselling_app:ledger_personal"))

        self.assertContains(expenses_resp, "Unique Grading Fee")
        self.assertContains(overview_resp, "25.00")
        self.assertNotContains(resale_resp, "Unique Grading Fee")
        self.assertNotContains(personal_resp, "Unique Grading Fee")

    def test_add_expense(self):
        response = self.client.post(reverse("albright_reselling_app:ledger_expenses"), {
            "add-date": "", "add-description": "Bubble wrap", "add-amount": "12.34", "add-category": "supplies",
            "add-notes": "", "add_entry": "1",
        })

        self.assertRedirects(response, reverse("albright_reselling_app:ledger_expenses"))
        expense = LedgerExpense.objects.get(owner=self.user, description="Bubble wrap")
        self.assertEqual(expense.amount, Decimal("12.34"))

    def test_edit_expense_via_formset(self):
        expense = _make_expense(self.user, description="Old Desc", amount="5.00")

        data = {
            "form-TOTAL_FORMS": "1", "form-INITIAL_FORMS": "1", "form-MIN_NUM_FORMS": "0",
            "form-MAX_NUM_FORMS": "1000",
            "form-0-id": str(expense.pk), "form-0-date": "", "form-0-description": "New Desc",
            "form-0-amount": "7.50", "form-0-category": "fees", "form-0-notes": "",
            "save_expenses": "1",
        }
        response = self.client.post(reverse("albright_reselling_app:ledger_expenses"), data)

        self.assertRedirects(response, reverse("albright_reselling_app:ledger_expenses"))
        expense.refresh_from_db()
        self.assertEqual(expense.description, "New Desc")
        self.assertEqual(expense.amount, Decimal("7.50"))
        self.assertEqual(expense.category, "fees")

    def test_delete_expense_via_formset(self):
        expense = _make_expense(self.user, description="To Delete", amount="5.00")

        data = {
            "form-TOTAL_FORMS": "1", "form-INITIAL_FORMS": "1", "form-MIN_NUM_FORMS": "0",
            "form-MAX_NUM_FORMS": "1000",
            "form-0-id": str(expense.pk), "form-0-date": "", "form-0-description": "To Delete",
            "form-0-amount": "5.00", "form-0-category": "supplies", "form-0-notes": "", "form-0-DELETE": "on",
            "save_expenses": "1",
        }
        self.client.post(reverse("albright_reselling_app:ledger_expenses"), data)

        self.assertFalse(LedgerExpense.objects.filter(pk=expense.pk).exists())

    def test_category_filter(self):
        _make_expense(self.user, description="Supplies Item", amount="5.00", category="supplies")
        _make_expense(self.user, description="Grading Item", amount="25.00", category="grading")

        response = self.client.get(reverse("albright_reselling_app:ledger_expenses"), {"category": "grading"})

        self.assertContains(response, "Grading Item")
        self.assertNotContains(response, "Supplies Item")

    def test_total_reflects_filtered_set(self):
        _make_expense(self.user, description="Supplies Item", amount="5.00", category="supplies")
        _make_expense(self.user, description="Grading Item", amount="25.00", category="grading")

        unfiltered = self.client.get(reverse("albright_reselling_app:ledger_expenses"))
        filtered = self.client.get(reverse("albright_reselling_app:ledger_expenses"), {"category": "grading"})

        self.assertContains(unfiltered, "$30.00")
        self.assertContains(filtered, "$25.00")

    def test_undated_shows_as_undated(self):
        _make_expense(self.user, description="No Date Item", amount="5.00", date_=None)

        response = self.client.get(reverse("albright_reselling_app:ledger_expenses"))

        self.assertContains(response, "Undated")


# ---------------------------------------------------------------------------
# Reselling-only metrics unaffected by expenses
# ---------------------------------------------------------------------------

class ResellingMetricsUnaffectedByExpensesTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="tester", password="pw-not-real-12345")

    def test_dashboard_tiles_unaffected(self):
        today = timezone.localdate()
        _make_entry(self.user, item="Resale Sold", cost="10.00", sold_for="40.00", status="sold_out",
                    sold_date=today)
        _make_expense(self.user, amount="1000.00", date_=today)

        tiles = ledger_metrics.dashboard_tiles(self.user)

        self.assertEqual(tiles["profit_this_month"], Decimal("30.00"))
        self.assertEqual(tiles["all_time_profit"], Decimal("30.00"))

    def test_scorecard_unaffected(self):
        _make_entry(self.user, item="Resale", cost="10.00", sold_for="20.00", status="sold_out")
        _make_expense(self.user, amount="500.00")

        context = ledger_metrics.scorecard_context(self.user)

        self.assertEqual(context["totals"]["total_profit"], Decimal("10.00"))

    def test_personal_items_context_unaffected(self):
        today = timezone.localdate()
        personal = _make_entry(self.user, item="Personal Sold", status="sold_out", ledger_type="personal")
        LedgerSale.objects.create(entry=personal, sale_date=today, sale_price=Decimal("20.00"))
        _make_expense(self.user, amount="999.00", date_=today)

        context = ledger_metrics.personal_items_context(self.user)

        self.assertEqual(context["tiles"]["revenue_this_month"], Decimal("20.00"))

    def test_net_cash_flow_resale_default_unaffected(self):
        _make_entry(self.user, item="Resale", cost="20.00", status="holding")
        _make_expense(self.user, amount="999.00")

        result = ledger_metrics.net_cash_flow(self.user)

        self.assertEqual(result["costs"], Decimal("20.00"))

    def test_resale_tab_tiles_dont_show_net_cash_flow_anymore(self):
        self.client.login(username="tester", password="pw-not-real-12345")
        response = self.client.get(reverse("albright_reselling_app:ledger_resale"))
        content = response.content.decode()
        self.assertNotIn('card__label">Net Cash Flow<', content)


# ---------------------------------------------------------------------------
# import_personal_ledger management command
# ---------------------------------------------------------------------------

REAL_SELLING_ROWS = [
    [date(2026, 5, 14), "Card A", 0, 2.66, 5.52, 6.00, 38.66, 24.48],
    [date(2026, 5, 18), "Card B", 0, 0, 2.33, 5.97, 12.97, 4.67],
    [date(2026, 5, 20), "Card C", 0, 2.95, 6.62, 5.97, 46.92, 31.38],
    [date(2026, 5, 21), "Card D", 0, 1.87, 3.65, 5.48, 24.52, 13.52],
    [date(2026, 5, 22), "Card E", 0, 0.94, 3.26, 5.62, 21.56, 11.74],
    [date(2026, 5, 24), "Card F", 0, 0.74, 1.91, 5.62, 11.36, 3.09],
    [date(2026, 6, 4), "Card G", 0, 3.26, 6.06, 5.48, 42.74, 27.94],
    [date(2026, 6, 5), "Card H", 0, 0, 7.89, 5.28, 50.28, 37.11],
    [date(2026, 6, 15), "Card I", 0, 1.66, 4.03, 5.72, 27.38, 15.97],
    [date(2026, 7, 15), "Pin A", 0, 1.89, 4.09, 5.24, 27.13, 15.91],
    [date(2026, 7, 26), "Pin B", 0, 1.45, 4.13, 5.97, 27.41, 15.86],
    [date(2026, 7, 26), "Pin C", 0, 3.64, 6.24, 5.83, 42.93, 27.22],
    [date(2026, 9, 15), "Card J", 0, 1.10, 3.64, 8.40, 24.49, 11.35],
    [date(2026, 9, 25), "Pin D", 0, 1.29, 3.08, 5.38, 19.67, 9.92],
    [date(2026, 10, 12), "Coin A", 0, 0, 9.68, 8.24, 68.24, 50.32],
]
REAL_EXPENSE_ROWS = [
    ["Bubble mailers + tape", 15.67],
    ["Mailer Envelopes", 18.26],
    ["Card fixing kit", 40.84],
    ["Penny Sleeves", 26.85],
    ["PSA Grading", None],
]


def _build_workbook_file(selling_rows, expense_rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Selling"
    ws.append(["Date", "Item", "Cost", "Sales Tax", "Fees", "Shipping", "Revenue", "Net Income"])
    for row in selling_rows:
        ws.append(row)
    ws2 = wb.create_sheet("Expenses")
    ws2.append(["Item", "Cost"])
    for row in expense_rows:
        ws2.append(row)
    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    tmp.close()
    wb.save(tmp.name)
    return tmp.name


class PlanSellingSheetUnitTests(TestCase):
    """Exercises _plan_selling_sheet()/_find_existing_sale() directly -
    these don't go through Command._check_expected_totals() at all, so
    small ad-hoc fixtures (not the full 15-row dataset) are enough."""

    def setUp(self):
        self.user = User.objects.create_user(username="tester", password="pw-not-real-12345")

    def _ws(self, rows):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Date", "Item", "Cost", "Sales Tax", "Fees", "Shipping", "Revenue", "Net Income"])
        for row in rows:
            ws.append(row)
        return ws

    def test_date_column_maps_to_sale_date(self):
        ws = self._ws([[date(2026, 5, 14), "Item A", 0, 0, 1.00, 1.00, 10.00, 8.00]])

        plan, mismatches = _plan_selling_sheet(ws, self.user)

        self.assertEqual(mismatches, [])
        self.assertEqual(plan[0].sale_date, date(2026, 5, 14))

    def test_blank_row_ignored(self):
        ws = self._ws([
            [None, None, None, None, None, None, None, None],
            [date(2026, 5, 14), "Item A", 0, 0, 1.00, 1.00, 10.00, 8.00],
        ])

        plan, mismatches = _plan_selling_sheet(ws, self.user)

        self.assertEqual(len(plan), 1)

    def test_sale_price_is_revenue_minus_sales_tax(self):
        ws = self._ws([[date(2026, 5, 14), "Item A", 0, 3.00, 1.00, 1.00, 20.00, 15.00]])

        plan, _mismatches = _plan_selling_sheet(ws, self.user)

        self.assertEqual(plan[0].sale_price, Decimal("17.00"))  # 20 - 3

    def test_mismatched_net_income_reported_and_excluded(self):
        ws = self._ws([[date(2026, 5, 14), "Item A", 0, 0, 1.00, 1.00, 10.00, 999.00]])

        plan, mismatches = _plan_selling_sheet(ws, self.user)

        self.assertEqual(plan, [])
        self.assertEqual(len(mismatches), 1)
        self.assertEqual(mismatches[0][0], "Item A")

    def test_matches_existing_entry_by_item_and_price(self):
        entry = LedgerEntry.objects.create(
            owner=self.user, item="Item A", ledger_type="personal", status="sold_out", cost=Decimal("0"),
        )
        LedgerSale.objects.create(entry=entry, sale_date=date(2020, 1, 1), sale_price=Decimal("8.00"))

        found_entry, found_sale = _find_existing_sale(self.user, "Item A", Decimal("8.00"))

        self.assertEqual(found_entry, entry)
        self.assertIsNotNone(found_sale)

    def test_no_match_when_price_differs(self):
        entry = LedgerEntry.objects.create(
            owner=self.user, item="Item A", ledger_type="personal", status="sold_out", cost=Decimal("0"),
        )
        LedgerSale.objects.create(entry=entry, sale_date=date(2020, 1, 1), sale_price=Decimal("8.00"))

        found_entry, _found_sale = _find_existing_sale(self.user, "Item A", Decimal("9.00"))

        self.assertIsNone(found_entry)


class ImportCommandFullRunTests(TestCase):
    """Full handle() runs, including Command._check_expected_totals() -
    the fixture here is numerically identical (same sums) to the real
    selling_ledger.xlsx snapshot the expected-totals constants are pinned
    to, with simplified item names."""

    def setUp(self):
        self.user = User.objects.create_user(username="tester", password="pw-not-real-12345")
        self.path = _build_workbook_file(REAL_SELLING_ROWS, REAL_EXPENSE_ROWS)

    def tearDown(self):
        os.unlink(self.path)

    def test_full_import_creates_expected_counts(self):
        out = StringIO()
        call_command("import_personal_ledger", self.path, stdout=out)

        self.assertEqual(LedgerEntry.objects.filter(ledger_type="personal").count(), 15)
        self.assertEqual(LedgerSale.objects.count(), 15)
        self.assertEqual(LedgerExpense.objects.count(), 4)
        self.assertIn("Expected totals check passed.", out.getvalue())

    def test_created_entries_have_no_buy_cost(self):
        call_command("import_personal_ledger", self.path, stdout=StringIO())

        for entry in LedgerEntry.objects.filter(ledger_type="personal"):
            self.assertEqual(entry.cost, Decimal("0"))
            self.assertEqual(entry.status, "sold_out")

    def test_sale_fields_mapped_correctly(self):
        call_command("import_personal_ledger", self.path, stdout=StringIO())

        entry = LedgerEntry.objects.get(item="Card A")
        sale = entry.sales.get()
        self.assertEqual(sale.sale_date, date(2026, 5, 14))
        self.assertEqual(sale.sale_price, Decimal("36.00"))  # 38.66 - 2.66
        self.assertEqual(sale.selling_fees, Decimal("5.52"))
        self.assertEqual(sale.outbound_shipping, Decimal("6.00"))
        self.assertEqual(sale.channel, "ebay")
        self.assertIsNotNone(sale.shipped_at)

    def test_idempotent_second_run_changes_nothing(self):
        call_command("import_personal_ledger", self.path, stdout=StringIO())
        call_command("import_personal_ledger", self.path, stdout=StringIO())

        self.assertEqual(LedgerEntry.objects.filter(ledger_type="personal").count(), 15)
        self.assertEqual(LedgerSale.objects.count(), 15)
        self.assertEqual(LedgerExpense.objects.count(), 4)

    def test_dry_run_saves_nothing(self):
        out = StringIO()
        call_command("import_personal_ledger", self.path, dry_run=True, stdout=out)

        self.assertEqual(LedgerEntry.objects.filter(ledger_type="personal").count(), 0)
        self.assertEqual(LedgerExpense.objects.count(), 0)
        self.assertIn("DRY RUN", out.getvalue())
        self.assertIn("Dry run - nothing was saved.", out.getvalue())

    def test_update_existing_entry_instead_of_duplicating(self):
        """Simulates a sheet imported once before dates were tracked: same
        item/price, a placeholder date, and a note flagging the date as
        unknown - re-importing now that the sheet has real dates must fix
        that one row in place, not create a 16th entry."""
        old_entry = LedgerEntry.objects.create(
            owner=self.user, item="Card A", ledger_type="personal", status="sold_out", cost=Decimal("0"),
        )
        old_sale = LedgerSale.objects.create(
            entry=old_entry, sale_date=date(2026, 1, 1), sale_price=Decimal("36.00"), channel="ebay",
            selling_fees=Decimal("5.52"), outbound_shipping=Decimal("6.00"),
            notes="Actual sale date unknown - estimated from eBay export order.",
        )

        call_command("import_personal_ledger", self.path, stdout=StringIO())

        self.assertEqual(LedgerEntry.objects.filter(ledger_type="personal", item="Card A").count(), 1)
        self.assertEqual(LedgerEntry.objects.filter(ledger_type="personal").count(), 15)  # not 16
        old_sale.refresh_from_db()
        self.assertEqual(old_sale.sale_date, date(2026, 5, 14))
        self.assertIsNotNone(old_sale.shipped_at)
        self.assertEqual(old_sale.notes, "")

    def test_unrelated_note_left_untouched_when_nothing_else_changes(self):
        """Re-running against an already-up-to-date row must be a no-op -
        including not touching a note that has nothing to do with the
        "unknown date" placeholder."""
        call_command("import_personal_ledger", self.path, stdout=StringIO())
        entry = LedgerEntry.objects.get(item="Card A")
        sale = entry.sales.get()
        sale.notes = "Shipped with extra padding."
        sale.save()

        call_command("import_personal_ledger", self.path, stdout=StringIO())

        sale.refresh_from_db()
        self.assertEqual(sale.notes, "Shipped with extra padding.")

    def test_mismatched_row_skipped_and_reported(self):
        bad_rows = REAL_SELLING_ROWS + [[date(2026, 10, 13), "Bad Row", 0, 0, 0, 0, 10.00, 999.00]]
        path = _build_workbook_file(bad_rows, REAL_EXPENSE_ROWS)
        out = StringIO()
        try:
            call_command("import_personal_ledger", path, stdout=out)

            self.assertFalse(LedgerEntry.objects.filter(item="Bad Row").exists())
            self.assertIn("MISMATCH", out.getvalue())
            self.assertEqual(LedgerEntry.objects.filter(ledger_type="personal").count(), 15)
        finally:
            os.unlink(path)

    def test_expense_with_no_amount_skipped_and_reported(self):
        out = StringIO()
        call_command("import_personal_ledger", self.path, stdout=out)

        self.assertFalse(LedgerExpense.objects.filter(description="PSA Grading").exists())
        self.assertIn("PSA Grading", out.getvalue())
        self.assertIn("SKIPPED", out.getvalue())

    def test_expense_category_guessed_from_description(self):
        rows = REAL_EXPENSE_ROWS + [["Grading submission fee", 12.00]]
        # Adjust the selling sheet count isn't needed - expected totals only
        # check the Selling sheet's own sums plus expense count/total/skips,
        # so adding a 5th real expense would fail that gate; test the
        # category guess directly against the plan instead.
        from albright_reselling_app.management.commands.import_personal_ledger import _plan_expenses_sheet
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Item", "Cost"])
        for row in rows:
            ws.append(row)

        plan, _skipped = _plan_expenses_sheet(ws, self.user)

        grading_row = next(r for r in plan if r.description == "Grading submission fee")
        self.assertEqual(grading_row.category, "grading")

    def test_expected_totals_mismatch_raises_and_saves_nothing(self):
        short_rows = REAL_SELLING_ROWS[:-1]  # only 14 valid rows - fails the expected count
        path = _build_workbook_file(short_rows, REAL_EXPENSE_ROWS)
        try:
            with self.assertRaises(CommandError):
                call_command("import_personal_ledger", path, stdout=StringIO())

            self.assertEqual(LedgerEntry.objects.filter(ledger_type="personal").count(), 0)
            self.assertEqual(LedgerExpense.objects.count(), 0)
        finally:
            os.unlink(path)

    def test_no_users_raises(self):
        User.objects.all().delete()
        with self.assertRaises(CommandError):
            call_command("import_personal_ledger", self.path, stdout=StringIO())

    def test_missing_file_raises(self):
        with self.assertRaises(CommandError):
            call_command("import_personal_ledger", "/no/such/file.xlsx", stdout=StringIO())

    def test_explicit_owner_argument(self):
        other = User.objects.create_user(username="other", password="pw-not-real-12345")
        call_command("import_personal_ledger", self.path, owner="other", stdout=StringIO())

        self.assertEqual(LedgerEntry.objects.filter(owner=other, ledger_type="personal").count(), 15)
        self.assertEqual(LedgerEntry.objects.filter(owner=self.user).count(), 0)
