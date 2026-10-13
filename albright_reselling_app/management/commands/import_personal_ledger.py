"""Imports selling_ledger.xlsx (project root, not checked in) into the
Personal Items ledger and the business-wide Expenses tab:

  - "Selling" sheet (Date, Item, Cost, Sales Tax, Fees, Shipping, Revenue,
    Net Income - extra columns to the right, the "Income:" total, are
    ignored): one sold-out personal LedgerEntry + one LedgerSale per row.
    sale_price is Revenue minus Sales Tax (eBay collects sales tax for the
    buyer - it was never the seller's income).
  - "Expenses" sheet (Item, Cost): one undated, category-guessed
    LedgerExpense per row. A row with no Cost is skipped and reported,
    never guessed at.

Idempotent: re-running with the same file changes nothing. A row matching
an existing personal entry by (item name, sale price) updates that sale's
date/shipped_at in place (and clears any leftover "actual sale date
unknown" note) instead of creating a duplicate - for a sheet imported once
before dates were tracked, then re-imported after they were added.
--dry-run reports what would change without saving anything.

Replaces this command's earlier --sale-date-only behavior now that the
sheet carries its own per-row Date column.
"""
from datetime import datetime, time
from decimal import ROUND_HALF_UP, Decimal

import openpyxl
from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError
from django.utils import timezone

from albright_reselling_app.models import LedgerEntry, LedgerExpense, LedgerSale

UNKNOWN_DATE_NOTE_SNIPPET = "actual sale date unknown"
MISMATCH_TOLERANCE = Decimal("0.01")

# Known-good totals for the current selling_ledger.xlsx snapshot - a hard
# sanity gate against a parsing/mapping mistake silently shipping wrong
# numbers. If the source spreadsheet's own data legitimately changes
# (new rows added, etc.), update these alongside it.
EXPECTED_SALE_COUNT = 15
EXPECTED_SALE_PRICE_TOTAL = Decimal("462.81")
EXPECTED_FEES_TOTAL = Decimal("72.13")
EXPECTED_SHIPPING_TOTAL = Decimal("90.20")
EXPECTED_NET_TOTAL = Decimal("300.48")
EXPECTED_EXPENSE_COUNT = 4
EXPECTED_EXPENSE_TOTAL = Decimal("101.62")
EXPECTED_EXPENSE_SKIPPED = ["PSA Grading"]


def _d(value):
    """Decimal, rounded to cents - openpyxl hands back floats, which carry
    binary-rounding noise ($24.480000000000004) that would otherwise
    accumulate across 15+ rows."""
    if value is None:
        return Decimal("0")
    return Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _as_date(value):
    if value is None:
        return None
    if hasattr(value, "date"):
        return value.date()
    return value


class _Row:
    def __init__(self, **kwargs):
        self.__dict__.update(kwargs)


def _plan_selling_sheet(ws, owner):
    """Reads every non-blank data row, validates it, and decides create/
    update/unchanged for each - no DB writes yet, so the expected-totals
    check below can run (and fail loudly) before anything is saved."""
    plan = []
    mismatches = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        date_val, item, cost, tax, fees, shipping, revenue, net_income = (row + (None,) * 8)[:8]
        if item is None or str(item).strip() == "":
            continue  # blank row

        item = str(item).strip()
        sale_date = _as_date(date_val)
        if sale_date is None:
            continue  # no date to import against - not a usable row
        cost = _d(cost)
        tax = _d(tax)
        fees = _d(fees)
        shipping = _d(shipping)
        revenue = _d(revenue)
        net_income = _d(net_income)

        sale_price = revenue - tax
        computed_net = revenue - tax - fees - shipping - cost
        if abs(computed_net - net_income) > MISMATCH_TOLERANCE:
            mismatches.append((item, net_income, computed_net))
            continue

        existing_entry, existing_sale = _find_existing_sale(owner, item, sale_price)
        plan.append(_Row(
            item=item, sale_date=sale_date, sale_price=sale_price, fees=fees, shipping=shipping,
            existing_entry=existing_entry, existing_sale=existing_sale,
        ))

    return plan, mismatches


def _find_existing_sale(owner, item_name, sale_price):
    entries = LedgerEntry.objects.filter(owner=owner, ledger_type="personal", item=item_name).prefetch_related(
        "sales"
    )
    for entry in entries:
        for sale in entry.sales.all():
            if sale.sale_price == sale_price:
                return entry, sale
    return None, None


def _plan_expenses_sheet(ws, owner):
    plan = []
    skipped = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        description, amount = (row + (None, None))[:2]
        if description is None or str(description).strip() == "":
            continue
        description = str(description).strip()
        if amount is None:
            skipped.append(description)
            continue

        amount = _d(amount)
        category = "grading" if "grading" in description.lower() else "supplies"
        existing = LedgerExpense.objects.filter(owner=owner, description=description, date__isnull=True).first()
        plan.append(_Row(description=description, amount=amount, category=category, existing=existing))

    return plan, skipped


class Command(BaseCommand):
    help = "Imports selling_ledger.xlsx's Selling + Expenses sheets (Personal Items ledger + business expenses)."

    def add_arguments(self, parser):
        parser.add_argument("path", help="Path to the .xlsx workbook.")
        parser.add_argument("--dry-run", action="store_true", help="Report what would change without saving.")
        parser.add_argument(
            "--owner", default=None,
            help="Username to own the imported rows. Defaults to the sole existing user.",
        )

    def handle(self, *args, **options):
        path = options["path"]
        dry_run = options["dry_run"]

        try:
            wb = openpyxl.load_workbook(path, data_only=True)
        except FileNotFoundError:
            raise CommandError(f"File not found: {path}")

        owner = self._resolve_owner(options["owner"])

        sale_plan, mismatches = _plan_selling_sheet(wb["Selling"], owner)
        expense_plan, expense_skipped = _plan_expenses_sheet(wb["Expenses"], owner)

        self._check_expected_totals(sale_plan, expense_plan, expense_skipped)

        self._apply_selling_plan(sale_plan, owner, dry_run)
        self._apply_expenses_plan(expense_plan, owner, dry_run)
        self._report(sale_plan, mismatches, expense_plan, expense_skipped, dry_run)

    def _resolve_owner(self, username):
        User = get_user_model()
        if username:
            try:
                return User.objects.get(username=username)
            except User.DoesNotExist:
                raise CommandError(f"No user named {username!r}.")
        users = list(User.objects.all()[:2])
        if not users:
            raise CommandError("No users exist to own the imported rows - create one first.")
        if len(users) > 1:
            raise CommandError("More than one user exists - pass --owner <username>.")
        return users[0]

    def _check_expected_totals(self, sale_plan, expense_plan, expense_skipped):
        sale_count = len(sale_plan)
        sale_price_total = sum((r.sale_price for r in sale_plan), Decimal("0"))
        fees_total = sum((r.fees for r in sale_plan), Decimal("0"))
        shipping_total = sum((r.shipping for r in sale_plan), Decimal("0"))
        net_total = sale_price_total - fees_total - shipping_total

        problems = []
        if sale_count != EXPECTED_SALE_COUNT:
            problems.append(f"sale count {sale_count} != expected {EXPECTED_SALE_COUNT}")
        if sale_price_total != EXPECTED_SALE_PRICE_TOTAL:
            problems.append(f"sale price total {sale_price_total} != expected {EXPECTED_SALE_PRICE_TOTAL}")
        if fees_total != EXPECTED_FEES_TOTAL:
            problems.append(f"fees total {fees_total} != expected {EXPECTED_FEES_TOTAL}")
        if shipping_total != EXPECTED_SHIPPING_TOTAL:
            problems.append(f"shipping total {shipping_total} != expected {EXPECTED_SHIPPING_TOTAL}")
        if net_total != EXPECTED_NET_TOTAL:
            problems.append(f"net total {net_total} != expected {EXPECTED_NET_TOTAL}")

        expense_count = len(expense_plan)
        expense_total = sum((r.amount for r in expense_plan), Decimal("0"))
        if expense_count != EXPECTED_EXPENSE_COUNT:
            problems.append(f"expense count {expense_count} != expected {EXPECTED_EXPENSE_COUNT}")
        if expense_total != EXPECTED_EXPENSE_TOTAL:
            problems.append(f"expense total {expense_total} != expected {EXPECTED_EXPENSE_TOTAL}")
        if sorted(expense_skipped) != sorted(EXPECTED_EXPENSE_SKIPPED):
            problems.append(f"expense skipped rows {expense_skipped} != expected {EXPECTED_EXPENSE_SKIPPED}")

        if problems:
            raise CommandError("Expected totals check failed:\n  " + "\n  ".join(problems))

    def _apply_selling_plan(self, plan, owner, dry_run):
        for row in plan:
            shipped_at = timezone.make_aware(datetime.combine(row.sale_date, time.min))
            if row.existing_sale is not None:
                sale = row.existing_sale
                changed = False
                if sale.sale_date != row.sale_date:
                    sale.sale_date = row.sale_date
                    changed = True
                if sale.shipped_at != shipped_at:
                    sale.shipped_at = shipped_at
                    changed = True
                if UNKNOWN_DATE_NOTE_SNIPPET in (sale.notes or "").lower():
                    sale.notes = ""
                    changed = True
                row.action = "update" if changed else "unchanged"
                if changed and not dry_run:
                    sale.save()
            else:
                row.action = "create"
                if not dry_run:
                    entry = LedgerEntry.objects.create(
                        owner=owner, item=row.item, ledger_type="personal", status="sold_out", cost=Decimal("0"),
                    )
                    LedgerSale.objects.create(
                        entry=entry, sale_date=row.sale_date, sale_price=row.sale_price, channel="ebay",
                        selling_fees=row.fees, outbound_shipping=row.shipping, shipped_at=shipped_at,
                    )

    def _apply_expenses_plan(self, plan, owner, dry_run):
        for row in plan:
            if row.existing is not None:
                expense = row.existing
                changed = expense.amount != row.amount or expense.category != row.category
                row.action = "update" if changed else "unchanged"
                if changed:
                    expense.amount = row.amount
                    expense.category = row.category
                    if not dry_run:
                        expense.save()
            else:
                row.action = "create"
                if not dry_run:
                    LedgerExpense.objects.create(
                        owner=owner, description=row.description, amount=row.amount, category=row.category,
                    )

    def _report(self, sale_plan, mismatches, expense_plan, expense_skipped, dry_run):
        prefix = "[DRY RUN] " if dry_run else ""
        self.stdout.write(f"{prefix}Selling sheet:")
        for row in sale_plan:
            self.stdout.write(f"  {row.action.upper():9s} {row.item!r}: {row.sale_date} for ${row.sale_price}")
        for item, net_income, computed_net in mismatches:
            self.stdout.write(self.style.WARNING(
                f"  MISMATCH  {item!r}: sheet net income ${net_income} != computed ${computed_net} - skipped"
            ))
        created = sum(1 for r in sale_plan if r.action == "create")
        updated = sum(1 for r in sale_plan if r.action == "update")
        unchanged = sum(1 for r in sale_plan if r.action == "unchanged")
        self.stdout.write(
            f"  {created} created, {updated} updated, {unchanged} unchanged, {len(mismatches)} mismatched/skipped"
        )

        self.stdout.write(f"{prefix}Expenses sheet:")
        for row in expense_plan:
            self.stdout.write(f"  {row.action.upper():9s} {row.description!r}: ${row.amount} ({row.category})")
        for description in expense_skipped:
            self.stdout.write(self.style.WARNING(f"  SKIPPED   {description!r}: no amount"))
        e_created = sum(1 for r in expense_plan if r.action == "create")
        e_updated = sum(1 for r in expense_plan if r.action == "update")
        e_unchanged = sum(1 for r in expense_plan if r.action == "unchanged")
        self.stdout.write(
            f"  {e_created} created, {e_updated} updated, {e_unchanged} unchanged, {len(expense_skipped)} skipped"
        )

        self.stdout.write(self.style.SUCCESS("Expected totals check passed."))
        if dry_run:
            self.stdout.write(self.style.WARNING("Dry run - nothing was saved."))
