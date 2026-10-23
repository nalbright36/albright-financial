"""My profile -> tailored resume and proposals, using ONLY my profile's facts.

How "never invent anything" is enforced (not just asked for):
  - Resume: the AI only picks and orders my skills, jobs, projects and
    bullet points BY INDEX. Employers, titles, dates and bullets in the
    output are copied from my profile - an index that doesn't exist is
    dropped, so an invented job can't appear.
  - Free text (the resume summary, the proposal): every sentence is checked
    against my profile and the posting. A sentence with a number, or a name
    (a capitalized word mid-sentence, e.g. an employer), that appears in
    neither is removed and listed under "removed claims" for me to judge.
  - Requirements I don't meet are listed (the AI's reading plus any
    required skill missing from my profile) - nothing is papered over.

Text formats for editing the profile (parse_* / format_*):
  Work history:  "Employer | Title | Start | End" then "- bullet" lines; blank line between jobs
  Projects:      "Name | link (optional)" then "- bullet" lines
  Education:     "School | Degree | Year" per line
  Certifications / Skills: one per line (skills may also be comma-separated)
  Links:         "Label | URL" per line
"""
import io
import re

from . import config
from .ai_tasks import run_json
from .models import ApplicationDraft
from .skills import normalize as normalize_skill

DEFAULT_RESUME_TEMPLATE = """{name}
{headline}
{links}

SUMMARY
{summary}

SKILLS
{skills}

EXPERIENCE
{experience}

PROJECTS
{projects}

EDUCATION
{education}

CERTIFICATIONS
{certifications}"""

DEFAULT_PROPOSAL_TEMPLATE = """Hi,

{proposal_body}

Best,
{name}
{links}"""

RESUME_PLACEHOLDERS = ("name", "headline", "summary", "skills", "experience", "projects", "education",
                       "certifications", "links")
PROPOSAL_PLACEHOLDERS = ("name", "headline", "proposal_body", "opportunity_title", "links")

# Capitalized words that are fine anywhere (greetings, pronouns, months, days...).
COMMON_WORDS = set("""I I'm I've I'd I'll Hi Hello Hey Dear Thanks Thank Best Regards Cheers Sincerely Kind Warm
Monday Tuesday Wednesday Thursday Friday Saturday Sunday January February March April May June July August
September October November December Jan Feb Mar Apr Jun Jul Aug Sep Sept Oct Nov Dec US USA UK EU
AM PM OK""".split())


# ---------------------------------------------------------------------------
# Profile text formats
# ---------------------------------------------------------------------------

def _split(line, parts):
    values = [v.strip() for v in line.split("|")]
    return (values + [""] * parts)[:parts]


def _blocks(text):
    """Header line + "- bullet" lines, blocks separated by blank lines or a new header."""
    blocks, current = [], None
    for raw in (text or "").splitlines():
        line = raw.strip()
        if not line:
            current = None
            continue
        if line[:1] in "-*•" and current is not None:
            bullet = line.lstrip("-*• ").strip()
            if bullet:
                current["bullets"].append(bullet)
            continue
        current = {"header": line, "bullets": []}
        blocks.append(current)
    return blocks


def parse_work_history(text):
    jobs = []
    for block in _blocks(text):
        employer, title, start, end = _split(block["header"], 4)
        jobs.append({"employer": employer, "title": title, "start": start, "end": end, "bullets": block["bullets"]})
    return jobs


def format_work_history(jobs):
    return "\n\n".join(" | ".join([j.get("employer", ""), j.get("title", ""), j.get("start", ""), j.get("end", "")])
                       + "".join(f"\n- {b}" for b in j.get("bullets") or []) for j in jobs or [])


def parse_projects(text):
    projects = []
    for block in _blocks(text):
        name, link = _split(block["header"], 2)
        projects.append({"name": name, "link": link, "bullets": block["bullets"]})
    return projects


def format_projects(projects):
    return "\n\n".join(" | ".join(filter(None, [p.get("name", ""), p.get("link", "")]))
                       + "".join(f"\n- {b}" for b in p.get("bullets") or []) for p in projects or [])


def parse_education(text):
    return [dict(zip(("school", "degree", "year"), _split(line, 3))) for line in (text or "").splitlines() if line.strip()]


def format_education(rows):
    return "\n".join(" | ".join(filter(None, [r.get("school", ""), r.get("degree", ""), r.get("year", "")])) for r in rows or [])


def parse_lines(text):
    return [line.strip() for line in (text or "").splitlines() if line.strip()]


def parse_skills(text):
    skills = []
    for part in re.split(r"[\n,]", text or ""):
        part = part.strip()
        if part and part.lower() not in (s.lower() for s in skills):
            skills.append(part[:60])
    return skills


def parse_links(text):
    return [dict(zip(("label", "url"), _split(line, 2))) for line in parse_lines(text)]


def format_links(links):
    return "\n".join(" | ".join(filter(None, [l.get("label", ""), l.get("url", "")])) for l in links or [])


# ---------------------------------------------------------------------------
# Fact checking free text
# ---------------------------------------------------------------------------

def profile_text(profile):
    """Every fact the drafts may use. Target roles are left out on purpose:
    they're what I'm looking for, not experience I have."""
    experience = f"{profile.years_experience} years of experience" if profile.years_experience else ""
    parts = [profile.name, profile.headline, profile.summary, " ".join(profile.skills or []), experience,
             ", ".join(profile.industries or []), profile.location,
             format_work_history(profile.work_history), format_projects(profile.projects),
             format_education(profile.education), "\n".join(profile.certifications or []), format_links(profile.links)]
    return "\n".join(p for p in parts if p)


def _numbers(text):
    return {re.sub(r"[,$%+]", "", n).rstrip(".") for n in re.findall(r"\$?\d[\d,]*(?:\.\d+)?%?\+?", text or "")}


def _words(text):
    return {w.lower() for w in re.findall(r"[A-Za-z][\w.+#&'-]*", text or "")}


def unsupported_claims(sentence, allowed_numbers, allowed_words):
    """Numbers and mid-sentence capitalized names in the sentence that
    appear in neither my profile nor the posting."""
    problems = [n for n in _numbers(sentence) if n not in allowed_numbers]
    tokens = re.findall(r"[A-Za-z][\w.+#&'-]*", sentence)
    for i, token in enumerate(tokens):
        clean = token.rstrip(".'")
        if i == 0 or not clean[:1].isupper() or clean in COMMON_WORDS:
            continue
        word = clean.lower()
        if word not in allowed_words and word.rstrip("s") not in allowed_words and word + "s" not in allowed_words:
            problems.append(clean)
    return problems


def check_text(text, profile, posting_text):
    """(text with unsupported sentences removed, [removed sentences with why])."""
    reference = f"{profile_text(profile)}\n{posting_text}"
    allowed_numbers, allowed_words = _numbers(reference), _words(reference)
    kept, removed = [], []
    for paragraph in (text or "").split("\n"):
        sentences = re.split(r"(?<=[.!?])\s+", paragraph.strip()) if paragraph.strip() else [""]
        kept_sentences = []
        for sentence in sentences:
            problems = unsupported_claims(sentence, allowed_numbers, allowed_words) if sentence else []
            if problems:
                removed.append(f"{sentence} (not in your profile: {', '.join(problems)})")
            else:
                kept_sentences.append(sentence)
        kept.append(" ".join(s for s in kept_sentences if s))
    return re.sub(r"\n{3,}", "\n\n", "\n".join(kept)).strip(), removed


# ---------------------------------------------------------------------------
# Requirements
# ---------------------------------------------------------------------------

def missing_skills(opp, profile):
    mine = {normalize_skill(s) for s in profile.skills or []}
    return [s for s in opp.skills_required or [] if normalize_skill(s) not in mine]


def _unmet(ai_list, opp, profile):
    unmet = [str(x).strip()[:200] for x in ai_list or [] if str(x).strip()]
    for skill in missing_skills(opp, profile):
        if not any(skill.lower() in u.lower() for u in unmet):
            unmet.append(f"{skill} (required skill not in your profile)")
    return unmet[:20]


def _posting(opp):
    return f"TITLE: {opp.title}\nREQUIRED SKILLS: {', '.join(opp.skills_required or []) or 'not listed'}\n\n{opp.raw_text}"


def _fill(template, values):
    """Replace only the known {placeholders}; any other braces are left alone."""
    return re.sub(r"\{(\w+)\}", lambda m: str(values[m.group(1)]) if m.group(1) in values else m.group(0), template)


def _links_line(profile):
    return " · ".join(f"{l.get('label')}: {l.get('url')}" if l.get("label") else l.get("url", "")
                      for l in profile.links or [] if l.get("url"))


# ---------------------------------------------------------------------------
# Tailored resume
# ---------------------------------------------------------------------------

RESUME_PROMPT = """You tailor a resume to a job posting using ONLY the candidate's profile below.
You may only SELECT and ORDER items by their index numbers - you cannot add or reword jobs, employers,
titles, dates, projects or bullet points. The summary may be rewritten (2-3 sentences) but must contain only
facts from the profile: no new employers, tools, numbers, years, or achievements.
Put what matters most for the posting first. Mirror the posting's key requirements in the summary only where
the profile genuinely supports them.

Return ONLY a JSON object:
{"summary": "...",
 "skills": [skill indexes, most relevant first],
 "work": [{"index": job index, "bullets": [bullet indexes, most relevant first]}],
 "projects": [{"index": project index, "bullets": [bullet indexes]}],
 "key_requirements": [the posting's key requirements],
 "unmet_requirements": [requirements the profile does NOT show]}"""


def _indexed_profile(profile):
    lines = [f"SUMMARY: {profile.summary}", "SKILLS:"]
    lines += [f"  s{i}: {s}" for i, s in enumerate(profile.skills or [])]
    lines.append("WORK HISTORY:")
    for i, job in enumerate(profile.work_history or []):
        lines.append(f"  job {i}: {job.get('title')} at {job.get('employer')} ({job.get('start')} - {job.get('end')})")
        lines += [f"    bullet {j}: {b}" for j, b in enumerate(job.get("bullets") or [])]
    lines.append("PROJECTS:")
    for i, project in enumerate(profile.projects or []):
        lines.append(f"  project {i}: {project.get('name')}")
        lines += [f"    bullet {j}: {b}" for j, b in enumerate(project.get("bullets") or [])]
    lines.append("EDUCATION: " + "; ".join(format_education([e]) for e in profile.education or []))
    lines.append("CERTIFICATIONS: " + "; ".join(profile.certifications or []))
    return "\n".join(lines)


def _index(value, size):
    try:
        i = int(str(value).lstrip("sS"))
    except (TypeError, ValueError):
        return None
    return i if 0 <= i < size else None


def _pick(entries, choices, source_list):
    """Entries chosen by index, with bullets chosen by index - copied from the profile, never from the AI."""
    picked, seen = [], set()
    for choice in choices if isinstance(choices, list) else []:
        index = _index(choice.get("index") if isinstance(choice, dict) else choice, len(source_list))
        if index is None or index in seen:
            continue
        seen.add(index)
        entry = dict(source_list[index])
        bullets = entry.get("bullets") or []
        chosen = [bullets[b] for b in (_index(x, len(bullets)) for x in (choice.get("bullets") or [] if isinstance(choice, dict) else []))
                  if b is not None]
        entry["bullets"] = list(dict.fromkeys(chosen)) or bullets
        picked.append(entry)
    for index, entry in enumerate(source_list):   # nothing true gets dropped - unchosen entries go last
        if index not in seen:
            picked.append(dict(entry))
    return picked


def render_resume(profile, structure):
    experience = "\n\n".join(
        f"{j['title']} - {j['employer']} ({j['start']} - {j['end']})".replace(" ()", "")
        + "".join(f"\n- {b}" for b in j.get("bullets") or []) for j in structure["work"])
    projects = "\n\n".join(p["name"] + (f" ({p['link']})" if p.get("link") else "")
                           + "".join(f"\n- {b}" for b in p.get("bullets") or []) for p in structure["projects"])
    template = profile.resume_template or DEFAULT_RESUME_TEMPLATE
    return _fill(template, {
        "name": profile.name, "headline": profile.headline, "summary": structure["summary"],
        "skills": ", ".join(structure["skills"]), "experience": experience, "projects": projects,
        "education": "\n".join(format_education([e]) for e in profile.education or []),
        "certifications": "\n".join(profile.certifications or []), "links": _links_line(profile),
    }).strip()


def tailor_resume(opp, profile, client=None):
    data, cost = run_json(RESUME_PROMPT, f"PROFILE:\n{_indexed_profile(profile)}\n\nPOSTING:\n{_posting(opp)}",
                          "resume", opportunity=opp, client=client, max_tokens=2500)
    skills = profile.skills or []
    chosen_skills = [skills[i] for i in (_index(x, len(skills)) for x in data.get("skills") or []) if i is not None]
    summary, removed = check_text(str(data.get("summary") or ""), profile, opp.raw_text)
    structure = {
        "summary": summary or profile.summary,
        "skills": list(dict.fromkeys(chosen_skills + skills)),
        "work": _pick(profile.work_history or [], data.get("work"), profile.work_history or []),
        "projects": _pick(profile.projects or [], data.get("projects"), profile.projects or []),
        "education": profile.education or [], "certifications": profile.certifications or [],
    }
    return ApplicationDraft.objects.create(
        opportunity=opp, kind="resume", text=render_resume(profile, structure), content=structure,
        key_requirements=[str(x)[:200] for x in data.get("key_requirements") or []][:15],
        unmet_requirements=_unmet(data.get("unmet_requirements"), opp, profile), removed_claims=removed, cost=cost)


# ---------------------------------------------------------------------------
# Proposal
# ---------------------------------------------------------------------------

PROPOSAL_PROMPT = """You write a short freelance proposal / cover letter body using ONLY facts from the candidate's
profile. Never invent experience, employers, clients, dates, skills, tools, numbers or results - if the profile
doesn't say it, don't write it. Open with one specific sentence about the client's problem as described in the
posting. Then connect it to the most relevant facts from the profile and mirror the posting's key requirements
only where the profile supports them. No greeting or sign-off (the template adds those). At most %d words.

Return ONLY a JSON object:
{"proposal_body": "...",
 "key_requirements": [the posting's key requirements],
 "unmet_requirements": [requirements the profile does NOT show]}"""


def word_count(text):
    return len(re.findall(r"\S+", text or ""))


def generate_proposal(opp, profile, client=None):
    max_words = int(config.get("proposal_max_words") or 200)
    data, cost = run_json(PROPOSAL_PROMPT % max_words,
                          f"PROFILE:\n{profile_text(profile)}\n\nPOSTING:\n{_posting(opp)}",
                          "proposal", opportunity=opp, client=client, max_tokens=1200)
    body, removed = check_text(str(data.get("proposal_body") or ""), profile, opp.raw_text)
    text = _fill(profile.proposal_template or DEFAULT_PROPOSAL_TEMPLATE, {
        "name": profile.name, "headline": profile.headline, "proposal_body": body,
        "opportunity_title": opp.title, "links": _links_line(profile)}).strip()
    return ApplicationDraft.objects.create(
        opportunity=opp, kind="proposal", text=text,
        content={"proposal_body": body, "word_count": word_count(body), "max_words": max_words},
        key_requirements=[str(x)[:200] for x in data.get("key_requirements") or []][:15],
        unmet_requirements=_unmet(data.get("unmet_requirements"), opp, profile), removed_claims=removed, cost=cost)


# ---------------------------------------------------------------------------
# .docx export
# ---------------------------------------------------------------------------

def resume_docx(profile, draft):
    """The tailored resume as a Word document (python-docx), from the draft's structure."""
    from docx import Document
    from docx.shared import Pt

    s = draft.content or {}
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.add_heading(profile.name or "Resume", level=0)
    if profile.headline:
        doc.add_paragraph(profile.headline)
    if _links_line(profile):
        doc.add_paragraph(_links_line(profile))
    if s.get("summary"):
        doc.add_heading("Summary", level=1)
        doc.add_paragraph(s["summary"])
    if s.get("skills"):
        doc.add_heading("Skills", level=1)
        doc.add_paragraph(", ".join(s["skills"]))
    if s.get("work"):
        doc.add_heading("Experience", level=1)
        for job in s["work"]:
            p = doc.add_paragraph()
            p.add_run(f"{job.get('title')} - {job.get('employer')}").bold = True
            if job.get("start") or job.get("end"):
                p.add_run(f"  ({job.get('start')} - {job.get('end')})")
            for bullet in job.get("bullets") or []:
                doc.add_paragraph(bullet, style="List Bullet")
    if s.get("projects"):
        doc.add_heading("Projects", level=1)
        for project in s["projects"]:
            p = doc.add_paragraph()
            p.add_run(project.get("name", "")).bold = True
            if project.get("link"):
                p.add_run(f"  {project['link']}")
            for bullet in project.get("bullets") or []:
                doc.add_paragraph(bullet, style="List Bullet")
    if s.get("education"):
        doc.add_heading("Education", level=1)
        for row in s["education"]:
            doc.add_paragraph(format_education([row]))
    if s.get("certifications"):
        doc.add_heading("Certifications", level=1)
        for cert in s["certifications"]:
            doc.add_paragraph(cert, style="List Bullet")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Prefill from my resume (python manage.py prefill_profile; Settings > Profile)
# ---------------------------------------------------------------------------

RESUME_PREFILL = {
    "target_roles": [
        "Data Analyst", "Business Intelligence Analyst/Developer", "Power BI Developer", "Tableau Developer",
        "Analytics Engineer", "Data Engineer", "Reporting Analyst", "Data Specialist", "Business Systems Analyst",
        "Salesforce Administrator", "Data Product Manager/Owner", "Automation/Integration Specialist",
        "Python/Django Developer (freelance/contract)",
    ],
    "skills": [
        "SQL", "Python", "Django", "PostgreSQL", "Power BI", "Tableau", "data modeling", "analytics engineering",
        "data products", "metrics/KPI design", "data quality", "business intelligence", "executive dashboards",
        "Salesforce (Certified Administrator)", "API integrations", "workflow automation", "ClickUp",
        "AI-assisted development", "project management", "business analysis", "team leadership",
    ],
    "years_experience": "8+",
    "industries": ["insurance/warranty", "healthcare/specialty pharma", "fintech (own project)"],
    "education": [{"school": "Illinois State University", "degree": "B.S. Finance", "year": "2018"}],
    "certifications": ["Salesforce Certified Administrator"],
    "location": "Tampa, FL",
}


def prefill_profile(profile):
    """Fills only what's empty (my edits are never overwritten); skills I
    already listed are kept, with the resume's added after them. Returns the
    names of the fields it changed."""
    changed = []
    for field, value in RESUME_PREFILL.items():
        current = getattr(profile, field)
        if field == "skills" and current:
            keys = {normalize_skill(s) for s in current}
            extra = [s for s in value if normalize_skill(s) not in keys]
            if extra:
                setattr(profile, field, list(current) + extra)
                changed.append(field)
        elif not current:
            setattr(profile, field, list(value) if isinstance(value, list) else value)
            changed.append(field)
    if changed:
        profile.save()
    return changed
