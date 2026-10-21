"""'Walk me through it' (GitHub read-only, no code execution), My profile,
tailored resumes and proposals that never add facts not in my profile,
.docx export, and paid tasks. GitHub and OpenAI are mocked."""
import io
import json
import re
from decimal import Decimal
from pathlib import Path
from unittest import mock

from django.conf import settings
from django.contrib.auth.models import User
from django.test import TestCase, override_settings
from django.urls import reverse

from albright_crypto_app import my_profile, walkthrough
from albright_crypto_app.models import (AIExtraction, ApplicationDraft, MyProfile, Opportunity, Source, TaskStep,
                                        Walkthrough)
from albright_crypto_app.pipeline import manual_source, rescore
from albright_crypto_app.tests.test_extraction import fake_client
from albright_crypto_app.tests.test_pipeline import make

CFG = {**settings.CRYPTO, "allowlist": ["github.com", "jobs.example.com"], "budget_usd": 0}
GH_GET = "albright_crypto_app.walkthrough.requests.get"
APP_DIR = Path(walkthrough.__file__).resolve().parent

ISSUE = {"number": 1, "title": "CSV parser bounty", "state": "open", "locked": False,
         "assignees": [{"login": "alice"}], "labels": [{"name": "bounty"}],
         "body": "The parser in src/parser/csv.py drops quoted newlines. Add a test in tests/test_csv.py.\n"
                 "Bounty: $200 in USDC, paid on merge by the maintainers."}
GITHUB = {
    "/issues/1": ISSUE,
    "/issues/1/comments": [{"user": {"login": "maint"}, "body": "Please comment before starting. See lib/util.py too."}],
    "/readme": "# Widget\n\npip install -e .\npytest",
    "/contents/.github/CONTRIBUTING.md": "Run `make test`. One PR per issue; reference the issue number.",
    "/languages": {"Python": 9000, "Shell": 1000},
    "/issues/1/timeline": [
        {"event": "cross-referenced", "source": {"issue": {"number": 7, "title": "Fix CSV newlines", "state": "open",
                                                           "pull_request": {}, "html_url": "https://github.com/acme/widget/pull/7"}}},
        {"event": "cross-referenced", "source": {"issue": {"number": 3, "title": "Old attempt", "state": "closed",
                                                           "pull_request": {}}}},
    ],
    "/contents/src/parser/csv.py": "def parse(text):\n    return text.split('\\n')\n",
    "/contents/lib/util.py": "import os\nos.system('rm -rf /')  # must never run\n",
}


class FakeGitHub:
    def __init__(self, routes=None):
        self.routes = routes or GITHUB
        self.urls = []

    def __call__(self, url, headers=None, params=None, timeout=None):
        self.urls.append(url)
        path = url.split("/repos/acme/widget", 1)[1]
        body = self.routes.get(path)
        resp = mock.Mock(status_code=200 if body is not None else 404)
        resp.json.return_value = body
        resp.text = body if isinstance(body, str) else json.dumps(body)
        return resp


GITHUB_PLAN = {
    "summary": "Fix quoted newlines in the CSV parser and add a test.",
    "claimable": "Assigned to alice and a PR is open - ask first.",
    "setup_steps": ["pip install -e .", "pytest"],
    "files_to_change": [{"path": "src/parser/csv.py", "why": "parser logic"}],
    "testing": ["make test"], "submit_pr": ["Reference #1 in the PR"],
    "payout_quote": "Bounty: $200 in USDC, paid on merge",
    "difficulty": "medium", "time_estimate_hours": 3, "time_estimate_note": "small parser fix",
    "questions": ["Is alice still working on this?"],
    "steps": [{"text": "Comment on the issue to ask if it's free"}, {"text": "Fix the parser"},
              {"text": "Add a test"}, {"text": "Open the PR"}],
}


def github_opp():
    return make(title="CSV parser bounty", url="https://github.com/acme/widget/issues/1")


@override_settings(CRYPTO=CFG)
class GitHubWalkthroughTests(TestCase):
    def generate(self, plan=GITHUB_PLAN, routes=None):
        fake = FakeGitHub(routes)
        with mock.patch(GH_GET, side_effect=fake):
            result = walkthrough.generate(github_opp() if not hasattr(self, "opp") else self.opp,
                                          client=fake_client(plan))
        return result, fake

    def setUp(self):
        self.opp = github_opp()

    def test_saves_a_plan_with_claimability_and_quoted_payout(self):
        w, fake = self.generate()
        self.assertEqual((w.opportunity, w.kind), (self.opp, "github"))
        self.assertGreater(w.cost, 0)
        self.assertEqual(AIExtraction.objects.get().purpose, "walkthrough")
        c = w.content
        self.assertIn("Already assigned to @alice - ask before starting.", c["warnings"])
        self.assertTrue(any("Open pull request #7" in x for x in c["warnings"]))
        self.assertFalse(any("#3" in x for x in c["warnings"]))          # closed PRs don't count
        self.assertEqual(c["payout_quote"], "Bounty: $200 in USDC, paid on merge")
        self.assertEqual((c["difficulty"], c["time_estimate_hours"]), ("medium", 3.0))
        titles = [s["title"] for s in c["sections"]]
        for title in ("What's being asked", "Still open / claimable?", "Local setup", "Likely files to change",
                      "How to test", "How to submit the PR"):
            self.assertIn(title, titles)
        self.assertEqual(c["questions"], ["Is alice still working on this?"])
        self.assertEqual(len(c["steps"]), 4)

    def test_reads_issue_comments_readme_contributing_languages_and_mentioned_files(self):
        w, fake = self.generate()
        paths = [u.split("/repos/acme/widget", 1)[1] for u in fake.urls]
        for path in ("/issues/1", "/issues/1/comments", "/readme", "/languages", "/issues/1/timeline",
                     "/contents/.github/CONTRIBUTING.md", "/contents/src/parser/csv.py", "/contents/lib/util.py",
                     "/contents/tests/test_csv.py"):
            self.assertIn(path, paths)
        repo = next(s for s in w.content["sections"] if s["title"] == "Repository")["items"]
        self.assertIn("src/parser/csv.py", repo[1])
        self.assertIn("Python", repo[0])

    def test_never_downloads_archives_clones_or_runs_code(self):
        with mock.patch("subprocess.Popen", side_effect=AssertionError("ran a process")), \
                mock.patch("subprocess.run", side_effect=AssertionError("ran a process")), \
                mock.patch("os.system", side_effect=AssertionError("ran a shell command")):
            w, fake = self.generate()
        self.assertTrue(fake.urls)
        for url in fake.urls:
            self.assertTrue(url.startswith("https://api.github.com/repos/acme/widget/"), url)
            self.assertNotRegex(url, r"tarball|zipball|archive|\.git$|/git/")
        # The file that would wipe a disk was only ever text handed to the AI.
        self.assertIn("lib/util.py", next(s for s in w.content["sections"] if s["title"] == "Repository")["items"][1])
        for name in ("walkthrough.py", "ai_tasks.py", "my_profile.py"):
            code = (APP_DIR / name).read_text(encoding="utf-8")
            self.assertNotRegex(code, r"\bsubprocess\b|os\.system|\bexec\(|\beval\(|importlib|tarfile|zipfile|"
                                      r"git clone|runpy|(?<!re\.)compile\(", name)

    def test_payout_quote_not_in_the_issue_is_dropped(self):
        w, _fake = self.generate({**GITHUB_PLAN, "payout_quote": "$5,000 guaranteed"})
        self.assertEqual(w.content["payout_quote"], "")

    def test_closed_issue_warns(self):
        routes = {**GITHUB, "/issues/1": {**ISSUE, "state": "closed", "assignees": []}, "/issues/1/timeline": []}
        w, _fake = self.generate(routes=routes)
        self.assertEqual(w.content["warnings"], ["This issue is closed."])

    def test_checklist_from_steps_once(self):
        TaskStep.objects.create(opportunity=self.opp, text="Fix the parser", order=1)
        w, _fake = self.generate()
        created = walkthrough.add_steps_to_checklist(w)
        self.assertEqual([s.text for s in created], ["Comment on the issue to ask if it's free", "Add a test",
                                                     "Open the PR"])   # "Fix the parser" already there
        self.assertEqual(walkthrough.add_steps_to_checklist(w), [])
        self.assertEqual(list(self.opp.task_steps.values_list("text", flat=True))[0], "Fix the parser")
        self.assertFalse(self.opp.task_steps.filter(automatable=True).exists())

    def test_mentioned_paths(self):
        self.assertEqual(walkthrough.mentioned_paths(
            "Bug in src/parser/csv.py; see https://github.com/a/b/blob/main/lib/x.ts and ../../etc/passwd.py"),
            ["lib/x.ts", "src/parser/csv.py"])
        self.assertEqual(walkthrough.github_issue_ref("https://github.com/acme/widget/issues/12"), ("acme", "widget", 12))
        self.assertIsNone(walkthrough.github_issue_ref("https://gitlab.com/acme/widget/issues/12"))


@override_settings(CRYPTO=CFG)
class WalkthroughPageTests(TestCase):
    def setUp(self):
        self.client.force_login(User.objects.create_user(username="tester", password="pw-not-real-12345"))
        self.opp = github_opp()

    def post(self, data, plan=GITHUB_PLAN):
        with mock.patch(GH_GET, side_effect=FakeGitHub()), \
                mock.patch("albright_crypto_app.ai_tasks.make_client", return_value=fake_client(plan)):
            return self.client.post(reverse("albright_crypto_app:opportunity", args=[self.opp.pk]), data)

    def test_button_saves_regenerates_and_reopens(self):
        page = self.client.get(reverse("albright_crypto_app:opportunity", args=[self.opp.pk]))
        self.assertContains(page, 'id="walkthrough-button">Walk me through it</button>')
        response = self.post({"action": "walkthrough"})
        self.assertRedirects(response, reverse("albright_crypto_app:opportunity", args=[self.opp.pk]) + "#walkthrough",
                             fetch_redirect_response=False)
        first = Walkthrough.objects.get()
        page = self.client.get(reverse("albright_crypto_app:opportunity", args=[self.opp.pk]))
        self.assertContains(page, "Bounty: $200 in USDC, paid on merge")
        self.assertContains(page, "Already assigned to @alice")
        self.assertContains(page, 'id="walkthrough-button">Regenerate</button>')

        self.post({"action": "walkthrough"}, {**GITHUB_PLAN, "summary": "Second take."})
        self.assertEqual(Walkthrough.objects.count(), 2)
        latest = self.client.get(reverse("albright_crypto_app:opportunity", args=[self.opp.pk]))
        self.assertContains(latest, "Second take.")
        older = self.client.get(reverse("albright_crypto_app:opportunity", args=[self.opp.pk]), {"walkthrough": first.pk})
        self.assertContains(older, "Fix quoted newlines")

    def test_one_click_checklist(self):
        self.post({"action": "walkthrough"})
        w = Walkthrough.objects.get()
        self.client.post(reverse("albright_crypto_app:opportunity", args=[self.opp.pk]),
                         {"action": "walkthrough_steps", "walkthrough": w.pk})
        self.assertEqual(self.opp.task_steps.count(), 4)

    def test_ai_limit_is_a_message(self):
        with override_settings(CRYPTO={**CFG, "ai_daily_limit": 0}):
            response = self.post({"action": "walkthrough"})
        self.assertFalse(Walkthrough.objects.exists())
        self.assertContains(self.client.get(response["Location"]), "Daily AI extraction limit reached")


@override_settings(CRYPTO=CFG)
class OtherWalkthroughTests(TestCase):
    def test_crypto_wallet_steps_are_always_manual(self):
        opp = make(title="Testnet quest", url="https://quests.example.org/q1", text="Bridge to the testnet and swap.")
        plan = {"summary": "Do a testnet quest.", "costs": "none stated", "difficulty": "easy",
                "steps": [{"text": "Read the quest page", "needs_wallet": False},
                          {"text": "Bridge test ETH to the testnet", "needs_wallet": False},   # AI forgot - caught
                          {"text": "Sign in with your wallet", "needs_wallet": True}]}
        w = walkthrough.generate(opp, client=fake_client(plan))
        self.assertEqual(w.kind, "crypto")
        self.assertEqual([s["needs_wallet"] for s in w.content["steps"]], [False, True, True])
        self.assertTrue(all(s["manual"] for s in w.content["steps"]))
        wallet_section = next(s for s in w.content["sections"] if s["title"].startswith("Wallet steps"))
        self.assertEqual(len(wallet_section["items"]), 2)
        created = walkthrough.add_steps_to_checklist(w)
        self.assertEqual(created[1].text, "Manual (wallet): Bridge test ETH to the testnet")
        self.assertEqual(w.content["warnings"], list(opp.safety_reasons))

    def test_income_walkthrough_has_red_flags_and_questions(self):
        opp = Opportunity.objects.create(source=manual_source("income"), opportunity_type="income", external_id="j1",
                                         title="Contract Django dev", url="https://jobs.example.com/1",
                                         raw_text="Contract Django role. Interview on Telegram only.")
        rescore(opp)
        plan = {"summary": "A Django contract.", "application_steps": ["Send a proposal"], "prepare": ["Portfolio link"],
                "red_flags": ["Interview only on a chat app"], "questions": ["What's the budget?"],
                "steps": [{"text": "Send a proposal"}], "difficulty": "easy", "time_estimate_hours": 1}
        w = walkthrough.generate(opp, client=fake_client(plan))
        self.assertEqual(w.kind, "income")
        titles = {s["title"]: s["items"] for s in w.content["sections"]}
        self.assertEqual(titles["Red flags noticed"], ["Interview only on a chat app"])
        self.assertEqual(titles["What to prepare"], ["Portfolio link"])
        self.assertEqual(w.content["questions"], ["What's the budget?"])


# ---------------------------------------------------------------------------
# My profile, tailored resume, proposal
# ---------------------------------------------------------------------------

PROFILE_POST = {
    "name": "Nate Albright", "headline": "Python & Django developer",
    "summary": "Python developer building Django apps and data pipelines for financial applications.",
    "skills": "Python, Django\nPostgreSQL\nExcel, VBA",
    "work_history": "Acme Corp | Software Developer | 2019-06 | 2024-03\n- Built Django APIs for billing\n"
                    "- Cut report time by 40% with VBA automation\n\n"
                    "Beta LLC | Data Analyst | 2016-01 | 2019-05\n- Built Excel models for 12 clients",
    "projects": "Lot scanner | https://github.com/example/scanner\n- Scrapes auctions and scores lots",
    "education": "State University | BS Finance | 2015",
    "certifications": "Microsoft Excel Expert",
    "links": "GitHub | https://github.com/example",
    "resume_template": "", "proposal_template": "",
}
POSTING = ("We need a Django developer to fix our billing API and automate Excel reports. "
           "Must know React and AWS. Fixed price $500.")


def income_opp():
    opp = Opportunity.objects.create(source=manual_source("income"), opportunity_type="income", external_id="p1",
                                     title="Django billing API + Excel automation", url="https://jobs.example.com/p1",
                                     raw_text=POSTING, skills_required=["Django", "React", "AWS", "Excel"])
    return rescore(opp)


@override_settings(CRYPTO=CFG)
class ProfileDraftTests(TestCase):
    def setUp(self):
        self.client.force_login(User.objects.create_user(username="tester", password="pw-not-real-12345"))
        self.assertEqual(self.client.post(reverse("albright_crypto_app:profile"), PROFILE_POST).status_code, 302)
        self.profile = MyProfile.load()
        self.opp = income_opp()

    def test_profile_is_parsed_into_structured_facts(self):
        p = self.profile
        self.assertEqual(p.skills, ["Python", "Django", "PostgreSQL", "Excel", "VBA"])
        self.assertEqual(p.work_history[0], {"employer": "Acme Corp", "title": "Software Developer", "start": "2019-06",
                                             "end": "2024-03", "bullets": ["Built Django APIs for billing",
                                                                           "Cut report time by 40% with VBA automation"]})
        self.assertEqual(p.projects[0]["link"], "https://github.com/example/scanner")
        self.assertEqual(p.education, [{"school": "State University", "degree": "BS Finance", "year": "2015"}])
        self.assertEqual(p.links, [{"label": "GitHub", "url": "https://github.com/example"}])
        page = self.client.get(reverse("albright_crypto_app:profile"))
        self.assertContains(page, "Acme Corp | Software Developer | 2019-06 | 2024-03")

    def test_profile_refuses_secrets(self):
        response = self.client.post(reverse("albright_crypto_app:profile"),
                                    {**PROFILE_POST, "summary": "seed phrase: " + " ".join(["apple"] * 12)})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(MyProfile.load().summary, PROFILE_POST["summary"])

    def test_tailored_resume_cannot_contain_an_invented_employer(self):
        ai = {"summary": "Django developer who built billing APIs. At Google I scaled payments to 10 million users.",
              "skills": ["s1", 3, 0, 42],                                  # 42 doesn't exist
              "work": [{"index": 0, "bullets": [0, 1]},
                       {"index": 5, "employer": "Google", "title": "Staff Engineer", "bullets": [0]}],  # invented job
              "projects": [{"index": 0, "bullets": [0]}],
              "key_requirements": ["Django", "billing API", "Excel automation"],
              "unmet_requirements": ["AWS"]}
        draft = my_profile.tailor_resume(self.opp, self.profile, client=fake_client(ai))

        employers = [job["employer"] for job in draft.content["work"]]
        self.assertEqual(employers, ["Acme Corp", "Beta LLC"])          # only profile employers, best first
        self.assertNotIn("Google", draft.text)
        self.assertNotIn("10 million", draft.text)
        self.assertEqual(draft.content["skills"][:3], ["Django", "Excel", "Python"])   # s1, 3, 0 (42 ignored)
        self.assertEqual(draft.content["summary"], "Django developer who built billing APIs.")
        self.assertTrue(any("Google" in r for r in draft.removed_claims))
        self.assertIn("AWS", draft.unmet_requirements)
        self.assertIn("React (required skill not in your profile)", draft.unmet_requirements)
        # Every fact in the output is in the profile.
        for job in draft.content["work"]:
            self.assertIn(job, self.profile.work_history)
        for skill in draft.content["skills"]:
            self.assertIn(skill, self.profile.skills)

    def test_proposal_drops_sentences_with_invented_facts(self):
        ai = {"proposal_body": "Your billing API needs fixing and your Excel reports need automating. "
                               "At Acme Corp I built Django APIs for billing. "
                               "At Google I led a team of 30 engineers. "
                               "I also cut report time by 40% with VBA automation.",
              "key_requirements": ["Django", "Excel"], "unmet_requirements": ["React", "AWS"]}
        draft = my_profile.generate_proposal(self.opp, self.profile, client=fake_client(ai))

        self.assertNotIn("Google", draft.text)
        self.assertNotIn("30", draft.text)
        self.assertIn("At Acme Corp I built Django APIs for billing.", draft.text)
        self.assertIn("40%", draft.text)                                # in my profile - kept
        self.assertEqual(len(draft.removed_claims), 1)
        self.assertIn("Google", draft.removed_claims[0])
        self.assertTrue(draft.text.startswith("Hi,\n\nYour billing API"))
        self.assertTrue(draft.text.rstrip().endswith("GitHub: https://github.com/example"))
        self.assertLessEqual(draft.content["word_count"], 200)
        self.assertEqual(draft.unmet_requirements, ["React", "AWS"])
        # Nothing left in the proposal that isn't in my profile or the posting.
        reference = my_profile.profile_text(self.profile) + "\n" + POSTING
        for sentence in re.split(r"(?<=[.!?])\s+", draft.content["proposal_body"]):
            self.assertEqual(my_profile.unsupported_claims(sentence, my_profile._numbers(reference),
                                                           my_profile._words(reference)), [], sentence)

    def test_prompt_says_under_200_words_by_default(self):
        client = fake_client({"proposal_body": "Your billing API needs fixing."})
        my_profile.generate_proposal(self.opp, self.profile, client=client)
        self.assertIn("At most 200 words", client.responses.create.call_args.kwargs["instructions"])

    def test_buttons_copy_download_docx_and_mark_sent(self):
        ai_resume = {"summary": "Django developer.", "skills": [1], "work": [{"index": 0, "bullets": [0]}],
                     "projects": [], "unmet_requirements": []}
        with mock.patch("albright_crypto_app.ai_tasks.make_client", return_value=fake_client(ai_resume)):
            self.client.post(reverse("albright_crypto_app:opportunity", args=[self.opp.pk]), {"action": "resume"})
        with mock.patch("albright_crypto_app.ai_tasks.make_client",
                        return_value=fake_client({"proposal_body": "Your billing API needs fixing."})):
            self.client.post(reverse("albright_crypto_app:opportunity", args=[self.opp.pk]), {"action": "proposal"})
        resume = ApplicationDraft.objects.get(kind="resume")
        proposal = ApplicationDraft.objects.get(kind="proposal")

        page = self.client.get(reverse("albright_crypto_app:opportunity", args=[self.opp.pk]))
        self.assertContains(page, f'data-copy-target="draft-text-{resume.pk}"')
        self.assertContains(page, f'id="draft-text-{proposal.pk}"')

        docx_response = self.client.get(reverse("albright_crypto_app:draft_download", args=[resume.pk]))
        self.assertEqual(docx_response["Content-Type"],
                         "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        from docx import Document
        text = "\n".join(p.text for p in Document(io.BytesIO(docx_response.content)).paragraphs)
        self.assertIn("Nate Albright", text)
        self.assertIn("Software Developer - Acme Corp", text)
        self.assertIn("Built Django APIs for billing", text)
        self.assertIn("State University", text)

        txt = self.client.get(reverse("albright_crypto_app:draft_download", args=[proposal.pk]))
        self.assertEqual(txt["Content-Type"], "text/plain; charset=utf-8")
        self.assertIn("Your billing API needs fixing.", txt.content.decode())

        self.client.post(reverse("albright_crypto_app:opportunity", args=[self.opp.pk]),
                         {"action": "draft_sent", "draft": proposal.pk})
        self.opp.refresh_from_db()
        proposal.refresh_from_db()
        self.assertEqual(self.opp.status, "applied")
        self.assertIsNotNone(proposal.sent_at)
        submitted = self.client.get(reverse("albright_crypto_app:income"), {"tab": "submitted"})
        self.assertIn(self.opp, list(submitted.context["opportunities"]))

    def test_empty_profile_sends_me_to_fill_it_in(self):
        MyProfile.objects.all().delete()
        with mock.patch("albright_crypto_app.ai_tasks.make_client") as make_client:
            response = self.client.post(reverse("albright_crypto_app:opportunity", args=[self.opp.pk]),
                                        {"action": "proposal"})
        self.assertRedirects(response, reverse("albright_crypto_app:profile"))
        make_client.assert_not_called()

    def test_custom_template_placeholders(self):
        self.profile.proposal_template = "Hello!\n{proposal_body}\n-- {name} ({headline}) re: {opportunity_title} {unknown}"
        self.profile.save()
        draft = my_profile.generate_proposal(self.opp, self.profile,
                                             client=fake_client({"proposal_body": "Your billing API needs fixing."}))
        self.assertEqual(draft.text, "Hello!\nYour billing API needs fixing.\n-- Nate Albright (Python & Django developer)"
                                     " re: Django billing API + Excel automation {unknown}")


@override_settings(CRYPTO=CFG)
class PaidTaskTests(TestCase):
    def setUp(self):
        self.client.force_login(User.objects.create_user(username="tester", password="pw-not-real-12345"))

    def test_track_a_paid_task_by_hand(self):
        with mock.patch("albright_crypto_app.pipeline.extract", side_effect=RuntimeError("no AI in this test")):
            self.client.post(reverse("albright_crypto_app:add"), {
                "opportunity_type": "income", "url": "https://jobs.example.com/annotate", "title": "AI training tasks",
                "text": "Rate chatbot answers. $20/hour.", "reward_type": "paid_task", "platform_name": "Annotation Co",
                "pay_rate_usd": "20", "application_status": "assessment"})
        opp = Opportunity.objects.get(title="AI training tasks")
        self.assertEqual((opp.reward_type, opp.platform_name, opp.pay_rate_usd, opp.pay_type, opp.application_status),
                         ("paid_task", "Annotation Co", Decimal("20.00"), "hourly", "assessment"))
        page = self.client.get(reverse("albright_crypto_app:opportunity", args=[opp.pk]))
        self.assertContains(page, 'id="paid-task"')
        self.assertContains(page, "Assessment")
        edit_fields = page.context["edit_form"].fields
        self.assertIn("platform_name", edit_fields)
        self.assertIn("paid_task", dict(edit_fields["reward_type"].choices))
